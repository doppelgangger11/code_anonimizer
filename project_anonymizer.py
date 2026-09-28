#!/usr/bin/env python3
"""
project_anonymizer.py — анонимизатор проектов (excel/csv + .py + .ipynb).

Что делает за один запуск:
  1. Анализирует Excel/CSV, находит текстовые колонки (эвристика по кардинальности).
  2. Даёт подтвердить/поправить список колонок (интерактивно или флагами).
  3. Строит ЕДИНЫЙ словарь замен  "оригинал -> токен"  (+ опционально переименование колонок).
  4. Создаёт копию проекта, где ТОТ ЖЕ словарь применён к:
       - Excel/CSV (значения; заголовки, если включено переименование),
       - .py            (только строки/комментарии — идентификаторы не трогаем),
       - .ipynb         (код-ячейки так же, markdown, outputs),
       - имена файлов/папок, а также .md/.txt/.sql/.yaml/.json...
     Хардкод вида df[df['client'] == 'ACME'] продолжает работать на анонимных данных.
  5. Сохраняет mapping.csv (ключ к де-анонимизации) ВНЕ выходной папки.

Использование:
  python project_anonymizer.py anonymize [ПАПКА] [-o OUT] [-m mapping.csv]
         [--columns "A,B"] [--exclude-columns "C"] [--rename-columns "X,Y"|all]
         [--threshold 0.5] [--clear-outputs] [--overwrite] [--yes]
  python project_anonymizer.py deanonymize ПАПКА_ANON -m mapping.csv -o OUT
Если ПАПКА не указана — спросит через input().
"""
from __future__ import annotations

import argparse
import bisect
import csv
import io
import json
import os
import re
import shutil
import sys
import tokenize
from collections import defaultdict
from pathlib import Path
from typing import Dict, Iterator, List, Optional, Tuple

import openpyxl

EXCEL_EXTS = {".xlsx", ".xlsm"}
CSV_EXTS = {".csv"}
TEXT_EXTS = {".md", ".txt", ".sql", ".yml", ".yaml", ".json", ".toml", ".cfg", ".ini"}
EXCLUDED_DIRS = {"backup", "archive", ".git", "__pycache__", ".ipynb_checkpoints",
                 ".venv", "venv", "node_modules"}
STOPVALUES = {"", "-", "—", "n/a", "na", "нет данных", "да", "нет", "yes", "no",
              "true", "false", "unknown", "неизвестно", "none", "nan", "null"}
LETTERS = "0-9A-Za-zА-Яа-яЁё"


# ----------------------------------------------------------------------------
# Вспомогательное
# ----------------------------------------------------------------------------

def sanitize_prefix(name: str) -> str:
    s = re.sub(rf"[^{LETTERS}]+", "_", str(name)).strip("_")
    return (s or "Value")[:30]


def is_stopvalue(v: str) -> bool:
    s = v.strip().lower()
    return (s in STOPVALUES or len(s) <= 1 or s.startswith("=")
            or re.fullmatch(r"[\d.,\-/: ]+", s) is not None)


def walk(root: Path, skip: Optional[Path] = None) -> Iterator[Path]:
    for dp, dns, fns in os.walk(root):
        dns[:] = sorted(d for d in dns if d.lower() not in EXCLUDED_DIRS
                        and (skip is None or (Path(dp) / d).resolve() != skip))
        for fn in sorted(fns):
            if not fn.startswith("~$"):
                yield Path(dp) / fn


def detect_header(rows: List[list]) -> int:
    for i, row in enumerate(rows):
        vals = [v for v in row if v not in (None, "")]
        if len(vals) >= 2 and sum(isinstance(v, str) for v in vals) >= 0.8 * len(vals):
            return i
    return 0


def read_csv(path: Path):
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1251"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("unknown encoding")
    try:
        delim = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|").delimiter
    except csv.Error:
        delim = ","
    return list(csv.reader(io.StringIO(text), delimiter=delim)), enc, delim


def iter_tables(path: Path):
    """yield (label, header, data_rows)"""
    if path.suffix.lower() in EXCEL_EXTS:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            for ws in wb.worksheets:
                rows = [list(r) for r in ws.iter_rows(values_only=True)]
                if not rows:
                    continue
                h = detect_header(rows[:30])
                header = [str(v).strip() if v not in (None, "") else f"col_{i}"
                          for i, v in enumerate(rows[h])]
                yield ws.title, header, rows[h + 1:]
        finally:
            wb.close()
    else:
        rows, _, _ = read_csv(path)
        if rows:
            h = detect_header(rows[:30])
            yield path.name, [c.strip() or f"col_{i}" for i, c in enumerate(rows[h])], rows[h + 1:]


# ----------------------------------------------------------------------------
# Анализ и построение словаря
# ----------------------------------------------------------------------------

class ColStats:
    def __init__(self):
        self.values, self.total = set(), 0


def collect_stats(tables_files: List[Path]) -> Dict[str, ColStats]:
    stats: Dict[str, ColStats] = defaultdict(ColStats)
    for p in tables_files:
        try:
            for _, header, rows in iter_tables(p):
                for c in header:
                    stats[c]  # зарегистрировать колонку (даже без текста)
                for row in rows:
                    for c, v in zip(header, row):
                        if isinstance(v, str):
                            st = stats[c]
                            st.total += 1
                            if not is_stopvalue(v):
                                st.values.add(v)
        except Exception as e:
            print(f"  [!] не смог прочитать {p}: {e}", file=sys.stderr)
    return stats


def choose_columns(stats, include, exclude, threshold) -> List[str]:
    out = []
    for c, st in stats.items():
        if c in exclude:
            continue
        if c in include or (st.total and len(st.values) >= 2
                            and len(st.values) / st.total >= threshold):
            out.append(c)
    return sorted(out)


def confirm_columns(stats, chosen: List[str]) -> List[str]:
    cols = sorted(c for c, s in stats.items() if s.total)
    sel = set(chosen)
    while True:
        print("\nТекстовые колонки ([x] = будет анонимизирована):")
        for i, c in enumerate(cols, 1):
            st = stats[c]
            print(f"  [{'x' if c in sel else ' '}] {i:>3}. {c!r}  ({len(st.values)} уник. / {st.total})")
        s = input("Номера для переключения (через пробел), Enter — принять: ").strip()
        if not s:
            return sorted(sel)
        for t in s.replace(",", " ").split():
            if t.isdigit() and 1 <= int(t) <= len(cols):
                sel ^= {cols[int(t) - 1]}


def build_value_map(files: List[Path], sensitive: List[str]):
    sens, vmap, counters = set(sensitive), {}, defaultdict(int)
    for p in files:
        try:
            for _, header, rows in iter_tables(p):
                idx = [i for i, c in enumerate(header) if c in sens]
                for row in rows:
                    for i in idx:
                        v = row[i] if i < len(row) else None
                        if isinstance(v, str) and not is_stopvalue(v) and v not in vmap:
                            pre = sanitize_prefix(header[i])
                            counters[pre] += 1
                            vmap[v] = f"{pre}_{counters[pre]:03d}"
        except Exception:
            continue
    return vmap


def build_column_map(cols: List[str]) -> Dict[str, str]:
    cols = [c for c in sorted(cols) if not re.fullmatch(r"col_\d+", c)]
    return {c: f"Column_{i:03d}" for i, c in enumerate(cols, 1)}


# ----------------------------------------------------------------------------
# Замена в тексте / коде
# ----------------------------------------------------------------------------

def make_pattern(mapping: Dict[str, str]) -> Optional[re.Pattern]:
    if not mapping:
        return None
    alt = "|".join(re.escape(k) for k in sorted(mapping, key=len, reverse=True))
    return re.compile(rf"(?<![{LETTERS}])(?:{alt})(?![{LETTERS}])")


def sub_text(text: str, mapping, pat, protected: Optional[List[Tuple[int, int]]] = None) -> str:
    if pat is None or not text:
        return text
    starts = [s for s, _ in protected] if protected else []

    def repl(m):
        if protected:
            i = bisect.bisect_right(starts, m.start()) - 1
            if i >= 0 and protected[i][1] > m.start():
                return m.group(0)
            if i + 1 < len(protected) and protected[i + 1][0] < m.end():
                return m.group(0)
        return mapping[m.group(0)]
    return pat.sub(repl, text)


def name_spans(code: str) -> List[Tuple[int, int]]:
    """Позиции идентификаторов (NAME-токенов): их менять нельзя.
    Строки IPython-магий (% и !) маскируем '#', сохраняя смещения."""
    parts = code.split("\n")
    lines = [l + "\n" for l in parts[:-1]] + [parts[-1]]
    masked = []
    for l in lines:
        s = l.lstrip()
        if s[:1] in ("%", "!"):
            i = len(l) - len(s)
            l = l[:i] + "#" + l[i + 1:]
        masked.append(l)
    offs = [0]
    for l in lines:
        offs.append(offs[-1] + len(l))
    spans = []
    for t in tokenize.generate_tokens(io.StringIO("".join(masked)).readline):
        if t.type == tokenize.NAME:
            (sr, sc), (er, ec) = t.start, t.end
            spans.append((offs[sr - 1] + sc, offs[er - 1] + ec))
    return sorted(spans)


def sub_code(code: str, mapping, pat) -> str:
    if pat is None:
        return code
    try:
        spans = name_spans(code)
    except (tokenize.TokenError, SyntaxError, IndentationError):
        return sub_text(code, mapping, pat)  # не разобрался — обычная замена по границам слов
    return sub_text(code, mapping, pat, spans)


# ----------------------------------------------------------------------------
# Обработчики файлов
# ----------------------------------------------------------------------------

class Transformer:
    def __init__(self, vmap, cmap, clear_outputs=False):
        self.vmap, self.cmap, self.clear_outputs = vmap, cmap, clear_outputs
        self.text_map = {**vmap, **cmap}
        self.pat = make_pattern(self.text_map)

    # --- excel ---
    def excel(self, src: Path, dst: Path):
        wb = openpyxl.load_workbook(src, keep_vba=src.suffix.lower() == ".xlsm")
        for ws in wb.worksheets:
            rows = list(ws.iter_rows())
            if not rows:
                continue
            h = detect_header([[c.value for c in r] for r in rows[:30]])
            for ri, row in enumerate(rows):
                for c in row:
                    v = c.value
                    if not isinstance(v, str) or v.startswith("="):
                        continue
                    new = (self.cmap.get(v.strip(), self.vmap.get(v)) if ri == h
                           else self.vmap.get(v))
                    if new is not None:
                        c.value = new
        wb.save(dst)

    # --- csv ---
    def csv(self, src: Path, dst: Path):
        rows, enc, delim = read_csv(src)
        if rows:
            h = detect_header(rows[:30])
            for ri, row in enumerate(rows):
                for ci, v in enumerate(row):
                    new = (self.cmap.get(v.strip(), self.vmap.get(v)) if ri == h
                           else self.vmap.get(v))
                    if new is not None:
                        row[ci] = new
        with open(dst, "w", encoding=enc, newline="") as f:
            csv.writer(f, delimiter=delim).writerows(rows)

    # --- py ---
    def py(self, src: Path, dst: Path):
        dst.write_text(sub_code(src.read_text(encoding="utf-8"), self.text_map, self.pat),
                       encoding="utf-8")

    def text(self, src: Path, dst: Path):
        dst.write_text(sub_text(src.read_text(encoding="utf-8"), self.text_map, self.pat),
                       encoding="utf-8")

    # --- ipynb ---
    def _src(self, source, fn):
        if isinstance(source, list):
            joined = fn("".join(source))
            return joined.splitlines(keepends=True)
        return fn(source) if isinstance(source, str) else source

    def ipynb(self, src: Path, dst: Path):
        nb = json.loads(src.read_text(encoding="utf-8"))
        code = lambda s: sub_code(s, self.text_map, self.pat)
        text = lambda s: sub_text(s, self.text_map, self.pat)
        for cell in nb.get("cells", []):
            fn = code if cell.get("cell_type") == "code" else text
            if "source" in cell:
                cell["source"] = self._src(cell["source"], fn)
            if cell.get("cell_type") != "code":
                continue
            if self.clear_outputs:
                cell["outputs"], cell["execution_count"] = [], None
                continue
            for out in cell.get("outputs", []) or []:
                if "text" in out:
                    out["text"] = self._src(out["text"], text)
                if "evalue" in out:
                    out["evalue"] = text(out["evalue"])
                if "traceback" in out:
                    out["traceback"] = [text(t) for t in out["traceback"]]
                data = out.get("data")
                if isinstance(data, dict):
                    for mime in list(data):
                        if mime.startswith("text/") or mime == "application/json":
                            if isinstance(data[mime], (str, list)):
                                data[mime] = self._src(data[mime], text)
        dst.write_text(json.dumps(nb, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")

    def rename(self, rel: Path) -> Path:
        return Path(*[sub_text(p, self.text_map, self.pat) for p in rel.parts])

    def process(self, src: Path, dst: Path):
        dst.parent.mkdir(parents=True, exist_ok=True)
        ext = src.suffix.lower()
        try:
            if ext in EXCEL_EXTS:
                self.excel(src, dst)
            elif ext in CSV_EXTS:
                self.csv(src, dst)
            elif ext == ".ipynb":
                self.ipynb(src, dst)
            elif ext == ".py":
                self.py(src, dst)
            elif ext in TEXT_EXTS:
                self.text(src, dst)
            else:
                shutil.copy2(src, dst)
        except Exception as e:
            print(f"  [!] {src}: {e} — файл скопирован без изменений", file=sys.stderr)
            shutil.copy2(src, dst)


def process_project(root: Path, out: Path, vmap, cmap, clear_outputs=False):
    tr = Transformer(vmap, cmap, clear_outputs)
    files = list(walk(root, skip=out.resolve()))
    for i, src in enumerate(files, 1):
        rel = src.relative_to(root)
        tr.process(src, out / tr.rename(rel))
        print(f"\r  обработано {i}/{len(files)}", end="", flush=True)
    print()
    return tr


# ----------------------------------------------------------------------------
# mapping.csv
# ----------------------------------------------------------------------------

def save_mapping(path: Path, vmap, cmap):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["kind", "original", "anonymized"])
        for k, v in cmap.items():
            w.writerow(["column", k, v])
        for k, v in vmap.items():
            w.writerow(["value", k, v])


def load_mapping(path: Path):
    vmap, cmap = {}, {}
    with open(path, encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            (cmap if r["kind"] == "column" else vmap)[r["original"]] = r["anonymized"]
    return vmap, cmap


def leftover_check(out: Path, vmap, cmap):
    """Ищем оригиналы, оставшиеся в .py/.ipynb (склейки строк, картинки и т.п.)."""
    keys = [k for k in {**vmap, **cmap} if len(k) >= 4]
    if not keys:
        return
    pat = re.compile("|".join(re.escape(k) for k in sorted(keys, key=len, reverse=True)))
    bad = []
    for p in walk(out):
        if p.suffix.lower() in (".py", ".ipynb"):
            try:
                m = pat.search(p.read_text(encoding="utf-8"))
            except Exception:
                continue
            if m:
                bad.append((p.relative_to(out), m.group(0)))
    if bad:
        print("\n[!] В этих файлах остались оригинальные значения — проверьте вручную")
        print("    (динамически собранные строки, значения в картинках-выводах и т.п.):")
        for p, k in bad:
            print(f"    - {p}  (например: {k!r})")
    else:
        print("\n✓ Проверка: оригинальных значений в .py/.ipynb не найдено.")


# ----------------------------------------------------------------------------
# Команды
# ----------------------------------------------------------------------------

def split_list(s: Optional[str]) -> set:
    return {x.strip() for x in s.split(",") if x.strip()} if s else set()


def cmd_anonymize(a):
    root = Path(a.root or input(">>> ")).expanduser().resolve()
    if not root.is_dir():
        sys.exit(f"Папка не найдена: {root}")
    out = Path(a.output).resolve() if a.output else root.parent / f"{root.name}_anonymized"
    map_file = Path(a.map_file).resolve() if a.map_file else root.parent / f"{root.name}_mapping.csv"
    if root in out.parents or out == root:
        sys.exit("Выходная папка не должна лежать внутри проекта.")
    if out.exists():
        if not a.overwrite:
            sys.exit(f"{out} уже существует (используйте --overwrite).")
        shutil.rmtree(out)

    all_files = list(walk(root))
    tables = [p for p in all_files if p.suffix.lower() in EXCEL_EXTS | CSV_EXTS]
    from collections import Counter
    cnt = Counter(p.suffix.lower() for p in all_files)
    print(f"Файлов: {len(all_files)} | excel: {sum(cnt[e] for e in EXCEL_EXTS)} "
          f"| csv: {cnt['.csv']} | py: {cnt['.py']} | ipynb: {cnt['.ipynb']}")

    stats = collect_stats(tables)
    chosen = choose_columns(stats, split_list(a.columns), split_list(a.exclude_columns), a.threshold)
    if not a.yes:
        chosen = confirm_columns(stats, chosen)
    print(f"\nАнонимизируемые колонки: {chosen}")
    vmap = build_value_map(tables, chosen)
    print(f"Уникальных значений для замены: {len(vmap)}")

    cmap = {}
    ren = a.rename_columns
    if ren is None and not a.yes:
        ren = input("Переименовать названия колонок? [Enter — нет / all / список через запятую]: ").strip() or None
    if ren:
        cols = list(stats) if ren.lower() == "all" else list(split_list(ren))
        cmap = build_column_map(cols)
        print(f"Колонок к переименованию: {len(cmap)}")

    save_mapping(map_file, vmap, cmap)  # сначала ключ — чтобы не потерять при сбое
    print("\nАнонимизация проекта...")
    process_project(root, out, vmap, cmap, a.clear_outputs)
    leftover_check(out, vmap, cmap)

    print("\nГотово.")
    print(f"  Проект:  {out}")
    print(f"  Ключ:    {map_file}   <- хранить ОТДЕЛЬНО, не публиковать")


def cmd_deanonymize(a):
    root, out = Path(a.root).resolve(), Path(a.output).resolve()
    vmap, cmap = load_mapping(Path(a.map_file))
    inv_v = {v: k for k, v in vmap.items()}
    inv_c = {v: k for k, v in cmap.items()}
    if out.exists() and not a.overwrite:
        sys.exit(f"{out} уже существует (используйте --overwrite).")
    if out.exists():
        shutil.rmtree(out)
    process_project(root, out, inv_v, inv_c)
    print(f"Восстановлено: {out}")


def main():
    p = argparse.ArgumentParser(description="Анонимизатор проектов (excel/csv + py + ipynb)")
    sub = p.add_subparsers(dest="cmd", required=True)

    a = sub.add_parser("anonymize")
    a.add_argument("root", nargs="?")
    a.add_argument("-o", "--output")
    a.add_argument("-m", "--map-file")
    a.add_argument("--columns", help="форсировать колонки (через запятую)")
    a.add_argument("--exclude-columns", help="исключить колонки (через запятую)")
    a.add_argument("--rename-columns", help="'all' или список названий колонок")
    a.add_argument("--threshold", type=float, default=0.5)
    a.add_argument("--clear-outputs", action="store_true", help="очистить outputs в ноутбуках")
    a.add_argument("--overwrite", action="store_true")
    a.add_argument("--yes", action="store_true", help="без интерактивных вопросов")
    a.set_defaults(func=cmd_anonymize)

    d = sub.add_parser("deanonymize")
    d.add_argument("root")
    d.add_argument("-m", "--map-file", required=True)
    d.add_argument("-o", "--output", required=True)
    d.add_argument("--overwrite", action="store_true")
    d.set_defaults(func=cmd_deanonymize)

    args = p.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
